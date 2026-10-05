import tempfile
from decimal import Decimal
from pathlib import Path

from django.test import SimpleTestCase
from openpyxl import Workbook

from core.adaptadores.plantilla.adaptador import AdaptadorPlantilla

ENCABEZADOS = [
    "NÚMERO DE DOCUMENTO", "TIPO DE COMPROBANTE", "CÓDIGO COMPROBANTE",
    "CUENTA CONTABLE", "DÉBITO O CRÉDITO", "NIT", "CÓDIGO DE LA CIUDAD",
    "DESCRIPCIÓN DE LA SECUENCIA", "SUCURSAL",
    "CUCON - CÓDIGO ÚNICO DE CONTRATO",
    "LÍNEA PRODUCTO", "GRUPO PRODUCTO", "CÓDIGO PRODUCTO", "CANTIDAD", "LOTE",
    "AÑO DEL VENCIMIENTO DEL LOTE",
]
BASE = ["100", "F", 1, "143505", "C", "900", "1", "Desc", "0"]
PRODUCTO = [150, 5, 65, 12, "L123", 2028]


class PlantillaCuconTest(SimpleTestCase):
    def _parsear(self, cucon):
        """cucon=None → formato viejo (la fila no trae la celda CUCON)."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Hoja1"
        for col, e in enumerate(ENCABEZADOS, start=1):
            ws.cell(row=5, column=col, value=e)
        fila = BASE + ([] if cucon is None else [cucon]) + PRODUCTO
        for col, v in enumerate(fila, start=1):
            ws.cell(row=6, column=col, value=v)
        with tempfile.TemporaryDirectory() as d:
            ruta = Path(d) / "corte.xlsx"
            wb.save(ruta)
            return AdaptadorPlantilla().parse(ruta)

    def _verificar(self, docs):
        linea = docs[0].lineas[0]
        self.assertEqual(linea.producto_codigo, "1500005000065")
        self.assertEqual(linea.cantidad_origen, Decimal("12"))
        self.assertEqual(linea.lote_raw, "L123")

    def test_formato_viejo_sin_celda_cucon(self):
        self._verificar(self._parsear(None))

    def test_formato_nuevo_cucon_texto(self):
        self._verificar(self._parsear("CT-0001"))

    def test_formato_nuevo_cucon_numerico(self):
        self._verificar(self._parsear(7))
