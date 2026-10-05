import io

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook

from cortes.models import Corte
from cortes.servicios.cargar import ErrorValidacionAdaptador, cargar_archivo

ENCABEZADOS = [
    "NÚMERO DE DOCUMENTO", "TIPO DE COMPROBANTE", "CÓDIGO COMPROBANTE",
    "CUENTA CONTABLE", "DÉBITO O CRÉDITO", "LÍNEA PRODUCTO", "GRUPO PRODUCTO",
    "CÓDIGO PRODUCTO", "CANTIDAD", "LOTE", "AÑO DEL VENCIMIENTO DEL LOTE",
    "NIT", "CÓDIGO DE LA CIUDAD", "DESCRIPCIÓN DE LA SECUENCIA", "SUCURSAL",
]


def _archivo(n_filas: int, anio) -> SimpleUploadedFile:
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    for col, e in enumerate(ENCABEZADOS, start=1):
        ws.cell(row=5, column=col, value=e)
    for i in range(n_filas):
        fila = [f"DOC{i}", "F", 1, "143505", "C", 150, 5, 5, 10, "L1", anio,
                "800000", "11001", "Desc", 0]
        for j, v in enumerate(fila, start=1):
            ws.cell(row=6 + i, column=j, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile("corte.xlsx", buf.getvalue())


class FormatoCambiadoTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("u", password="x")

    def test_columnas_desalineadas_se_rechazan(self):
        # año de vencimiento = 0 en todas las filas: las columnas no cuadran
        with self.assertRaises(ErrorValidacionAdaptador) as ctx:
            cargar_archivo(_archivo(6, 0), self.user, "PLANTILLA", 1)
        self.assertIn("formato", str(ctx.exception))
        self.assertEqual(Corte.objects.count(), 0)

    def test_todo_sin_maestro_se_rechaza_y_no_deja_corte(self):
        # catálogo vacío: las 12 líneas quedan sin maestro
        with self.assertRaises(ErrorValidacionAdaptador):
            cargar_archivo(_archivo(12, 2028), self.user, "PLANTILLA", 1)
        self.assertEqual(Corte.objects.count(), 0)

    def test_corte_chico_sin_maestro_si_se_acepta(self):
        corte, resultado = cargar_archivo(_archivo(3, 2028), self.user, "PLANTILLA", 1)
        self.assertEqual(corte.estado, "en_revision")
        self.assertEqual(resultado.lineas_sin_maestro, 3)
